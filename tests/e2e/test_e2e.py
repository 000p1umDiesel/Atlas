"""Сквозной сценарий на настоящих сервисах.

Нужны Docker (testcontainers: Neo4j + Qdrant из `tests/conftest.py`), локальный Ollama с
`embeddinggemma-2:740m` и LLM-эндпоинт из `config.yaml` с `LLM_API_KEY` в `.env`.
Запуск: `uv run pytest -m e2e -s`.
"""

import shutil
from pathlib import Path

import pytest

from mnogobase.app import build_app
from mnogobase.config import load_settings
from mnogobase.ids import file_doc_id, normalize_name
from mnogobase.retrieval import Mode, build_retrievers
from mnogobase.retrieval.answer import Answerer
from mnogobase.retrieval.compare import run_mode
from mnogobase.stores.qdrant_store import QdrantStore

pytestmark = [
    pytest.mark.e2e,
    # выбрасывается внутри обработки OCR-опций самого Docling, а не mnogobase
    pytest.mark.filterwarnings("ignore:`force_full_page_ocr` is deprecated:DeprecationWarning"),
    pytest.mark.filterwarnings(
        "ignore:deprecated:DeprecationWarning:docling.models.stages.ocr.rapid_ocr_model"
    ),
]
ROOT = Path(__file__).parents[2]
FIXTURES = ROOT / "tests" / "fixtures"
QUESTION = "What is the Transformer and who proposed it?"


def make_corpus(folder: Path) -> None:
    from docx import Document
    from reportlab.pdfgen import canvas

    folder.mkdir()
    for name in ("attention_en.md", "vnimanie_ru.md"):
        shutil.copy(FIXTURES / name, folder / name)
    docx = Document()
    docx.add_heading("Transformer architecture notes", 0)
    docx.add_paragraph(
        "The Transformer replaced recurrent networks in machine translation. "
        "Its attention mechanism uses softmax to weight the values."
    )
    docx.save(folder / "notes.docx")
    pdf = canvas.Canvas(str(folder / "rag.pdf"))
    lines = [
        "Retrieval-Augmented Generation",
        "",
        "Retrieval-augmented generation (RAG) combines a retriever with a language model.",
        "Documents are embedded with a Transformer encoder and stored in a vector database.",
        "At question time the most similar passages are added to the prompt.",
    ]
    text = pdf.beginText(72, 750)
    for line in lines:
        text.textLine(line)
    pdf.drawText(text)
    pdf.save()


def entity_ids_of(graph, doc_id: str) -> set[str]:
    return {eid for eids in graph.chunk_entity_ids(doc_id).values() for eid in eids}


async def test_full_flow(graph, qdrant_client, tmp_path, monkeypatch):
    monkeypatch.chdir(ROOT)  # load_settings читает `.env` (LLM_API_KEY) из CWD
    corpus = tmp_path / "corpus"
    make_corpus(corpus)
    base = load_settings(ROOT / "config.yaml")
    settings = base.model_copy(
        update={
            "data_dir": tmp_path / ".mb",
            "logs_dir": tmp_path / "logs",
            "runs_dir": tmp_path / "runs",
            "wiki": base.wiki.model_copy(update={"dir": tmp_path / "wiki"}),
        }
    )
    vectors = QdrantStore(qdrant_client, "e2e_", settings.embedder.dim)
    app = build_app(settings, vectors=vectors, graph=graph)
    try:
        report = await app.pipeline.ingest([corpus])
        assert report.failed == {}, report.failed
        assert len(report.processed) == 4
        assert report.wiki is not None and report.wiki.failed == [], report.wiki
        assert report.wiki.created

        names = [e.name for e in graph.entities()]
        transformers = [e for e in graph.entities() if normalize_name(e.name) == "transformer"]
        assert transformers, names
        assert max(e.mention_count for e in transformers) >= 2, names

        # межъязыковое слияние: упоминания на RU и EN сводятся к одной сущности
        en_ids = entity_ids_of(graph, file_doc_id(corpus / "attention_en.md"))
        ru_ids = entity_ids_of(graph, file_doc_id(corpus / "vnimanie_ru.md"))
        assert en_ids and ru_ids
        shared = en_ids & ru_ids
        assert shared, names
        assert any(t.entity_id in shared for t in transformers), (
            [graph.get_entity(eid).name for eid in shared],
            names,
        )

        pages = list((tmp_path / "wiki" / "entities").glob("*.md"))
        assert pages
        assert any("[^" in p.read_text(encoding="utf-8") for p in pages)
        assert (tmp_path / "wiki" / "index.md").is_file()

        retrievers = build_retrievers(vectors, graph, app.embedder, app.sparse, settings)
        answerer = Answerer(app.llm, settings.retrieval.context_tokens)
        for mode in Mode:
            answer = await run_mode(
                QUESTION, mode.value, retrievers[mode.value], answerer, settings.retrieval.k
            )
            assert answer.text.strip(), mode
            assert answer.sources, mode
            assert any(s.cited for s in answer.sources), (mode, answer.text)

        again = await app.pipeline.ingest([corpus])
        assert again.processed == [] and len(again.skipped) == 4, again
        assert again.failed == {}
    finally:
        app.registry.close()  # драйвер Neo4j закрывает фикстура `graph`
